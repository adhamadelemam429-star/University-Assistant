from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

import fitz
from docx import Document

from app.config import settings
from app.models.schemas import ExtractedDocument


def _normalize_text(raw_text: str) -> str:
    text = raw_text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'\s{2,}', ' ', text)
    return text.strip()


def extract_text_from_file(file_path: str | Path) -> ExtractedDocument:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f'File not found: {path}')

    suffix = path.suffix.lower()
    if suffix not in settings.allowed_extensions:
        raise ValueError(f'Unsupported file type: {suffix}')

    try:
        if suffix == '.txt':
            text = path.read_text(encoding='utf-8', errors='ignore')
            cleaned = _normalize_text(text)
            return ExtractedDocument(file_name=path.name, text=cleaned, is_empty=not bool(cleaned.strip()))

        if suffix == '.csv':
            rows = []
            with path.open('r', encoding='utf-8-sig', errors='ignore', newline='') as csv_file:
                reader = csv.reader(csv_file)
                for row in reader:
                    rows.append(' | '.join(part.strip() for part in row if part and part.strip()))
            text = '\n'.join(rows)
            cleaned = _normalize_text(text)
            return ExtractedDocument(file_name=path.name, text=cleaned, is_empty=not bool(cleaned.strip()))

        if suffix == '.docx':
            doc = Document(str(path))
            paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
            cleaned = _normalize_text('\n'.join(paragraphs))
            return ExtractedDocument(file_name=path.name, text=cleaned, is_empty=not bool(cleaned.strip()))

        if suffix == '.pdf':
            document = fitz.open(str(path))
            page_texts: list[str] = []
            for page_number in range(document.page_count):
                page = document[page_number]
                page_text = page.get_text('text', sort=True)
                page_texts.append(page_text)
            text = '\n\n'.join(page_texts).strip()
            cleaned = _normalize_text(text)
            if not cleaned:
                return ExtractedDocument(
                    file_name=path.name,
                    text='',
                    is_empty=True,
                    is_scanned_pdf=True,
                    warning='No selectable text was found in this PDF. It may be a scanned document without OCR support.',
                )
            return ExtractedDocument(file_name=path.name, text=cleaned, page_number=1, is_empty=False)
    except Exception as exc:  # pragma: no cover - defensive path
        return ExtractedDocument(file_name=path.name, text='', is_empty=True, warning=str(exc))


def load_documents_from_directory(directory: str | Path) -> list[ExtractedDocument]:
    folder = Path(directory)
    results: list[ExtractedDocument] = []
    if not folder.exists():
        return results

    for file_path in sorted(folder.iterdir()):
        if file_path.is_file() and file_path.suffix.lower() in settings.allowed_extensions:
            try:
                results.append(extract_text_from_file(file_path))
            except Exception:
                continue
    return results
